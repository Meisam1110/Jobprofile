"""Export a job profile to Word (official Irancell template layout), PDF,
Excel and JSON."""
from __future__ import annotations

import io
import json

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF

from .. import models

BRAND_YELLOW = "FFCB05"
DARK = "1A1A1A"

TEMPLATE_SECTIONS = [
    ("Mission", "mission"),
    ("Context", "context"),
    ("Roles & Responsibilities", "responsibilities_text"),
    ("Collaboration", "collaboration_text"),
    ("Authorities", "authorities_text"),
]
REQUIREMENT_SECTIONS = [
    ("Education", "education_text"),
    ("Experience", "experience_text"),
]
TAIL_SECTIONS = [
    ("Trainings", "trainings_text"),
    ("General Working Condition", "working_conditions"),
    ("Performance Standards", "performance_standards"),
]


def _grade(profile: models.JobProfile) -> str:
    return profile.grade.code if profile.grade else (profile.org_level or "")


def profile_to_dict(profile: models.JobProfile) -> dict:
    return {
        "id": profile.id,
        "job_title": profile.job_title,
        "job_code": profile.job_code,
        "location": profile.location,
        "reports_to": profile.reports_to,
        "direct_reports": profile.direct_reports,
        "level": _grade(profile),
        "band": profile.band,
        "division": profile.division.name if profile.division else "",
        "department": profile.department.name if profile.department else "",
        "role_family": profile.role_family.name if profile.role_family else "",
        "job_family": profile.job_family.name if profile.job_family else "",
        "status": profile.status,
        "version": profile.version,
        "document_version": profile.document_version,
        "approved_date": profile.approved_date.isoformat() if profile.approved_date else None,
        "mission": profile.mission,
        "context": profile.context,
        "responsibilities": [r.text for r in profile.responsibilities],
        "collaboration": profile.collaboration,
        "authorities": profile.authorities_text,
        "education": profile.education_text,
        "experience": profile.experience_text,
        "technical_competencies": [
            {"name": pc.competency.name, "category": pc.competency.category, "level": pc.required_level}
            for pc in profile.profile_competencies
            if pc.competency and pc.competency.comp_type == "Technical"
        ],
        "behavioral_competencies": [
            pc.competency.name
            for pc in profile.profile_competencies
            if pc.competency and pc.competency.comp_type == "Behavioral"
        ],
        "trainings": profile.trainings_text,
        "working_conditions": profile.working_conditions,
        "performance_standards": profile.performance_standards,
        "signoff": profile.signoff,
        "kpis": [
            {"name": k.name, "type": k.kpi_type, "weight": k.weight, "measurement": k.measurement_method}
            for k in profile.kpis
        ],
    }


def to_json(profile: models.JobProfile) -> bytes:
    return json.dumps(profile_to_dict(profile), indent=2, ensure_ascii=False).encode("utf-8")


# ------------------------------------------------------------------- Word
def _docx_section(document: Document, heading: str, body: str):
    table = document.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    head_cell = table.rows[0].cells[0]
    run = head_cell.paragraphs[0].add_run(heading)
    run.bold = True
    run.font.size = Pt(11)
    body_cell = table.rows[1].cells[0]
    body_cell.paragraphs[0].add_run(body or "—").font.size = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def to_docx(profile: models.JobProfile) -> bytes:
    """Replicates the official template: header block table, boxed sections,
    Job Requirements + Competencies groups, Signoff grid."""
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("JOB PROFILE — MTN IRANCELL")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(DARK)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Organization Architecture, HR").font.size = Pt(9)

    header = document.add_table(rows=2, cols=3)
    header.style = "Table Grid"
    cells = [
        ("Title of Position:", profile.job_title),
        ("Location of the Job:", profile.location),
        ("Reports to:", profile.reports_to),
        ("Number of subordinates:", str(profile.direct_reports)),
        ("Level:", _grade(profile)),
        ("Job Code:", profile.job_code),
    ]
    for idx, (label, value) in enumerate(cells):
        cell = header.rows[idx // 3].cells[idx % 3]
        p = cell.paragraphs[0]
        p.add_run(label + "\n").bold = True
        p.add_run(value or "—").font.size = Pt(10)
    document.add_paragraph()

    for heading, attr in TEMPLATE_SECTIONS:
        _docx_section(document, heading, getattr(profile, attr))

    _docx_section(document, "Job Requirements", "")
    for heading, attr in REQUIREMENT_SECTIONS:
        _docx_section(document, heading, getattr(profile, attr))

    _docx_section(document, "Competencies", "")
    tech = [pc for pc in profile.profile_competencies if pc.competency and pc.competency.comp_type == "Technical"]
    if tech:
        table = document.add_table(rows=1 + len(tech), cols=3)
        table.style = "Table Grid"
        for col, text in enumerate(["Title", "Type", "Level"]):
            table.rows[0].cells[col].paragraphs[0].add_run(text).bold = True
        for row, pc in enumerate(tech, start=1):
            table.rows[row].cells[0].text = pc.competency.name
            table.rows[row].cells[1].text = pc.competency.category or "Technical"
            table.rows[row].cells[2].text = pc.required_level
        document.add_paragraph()
    behavioral = [pc.competency.name for pc in profile.profile_competencies if pc.competency and pc.competency.comp_type == "Behavioral"]
    _docx_section(document, "Behavioral Competencies", "\n".join(f"• {b}" for b in behavioral) or profile.behavioral_text)

    for heading, attr in TAIL_SECTIONS:
        _docx_section(document, heading, getattr(profile, attr))

    signoff = document.add_table(rows=3, cols=3)
    signoff.style = "Table Grid"
    signoff.rows[0].cells[0].merge(signoff.rows[0].cells[2])
    signoff.rows[0].cells[0].paragraphs[0].add_run("Signoff").bold = True
    labels = ["Line Manager:", "Functional Manager:", "OA:", "HoD:", "CHRO: (if required)", "CEO/COO: (if required)"]
    for idx, label in enumerate(labels):
        signoff.rows[1 + idx // 3].cells[idx % 3].text = label

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------------- PDF
def _pdf_text(text: str) -> str:
    safe = (text or "—").encode("latin-1", "replace").decode("latin-1")
    # fpdf cannot wrap single words longer than the line; soft-break them
    return " ".join(
        word if len(word) <= 50 else " ".join(word[i : i + 50] for i in range(0, len(word), 50))
        for word in safe.split(" ")
    )


def to_pdf(profile: models.JobProfile) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_fill_color(255, 203, 5)  # MTN yellow
    pdf.rect(0, 0, 210, 16, style="F")
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_y(4)
    pdf.cell(0, 8, _pdf_text(f"JOB PROFILE - MTN IRANCELL"), align="C")
    pdf.ln(14)

    pdf.set_font("Helvetica", "B", 14)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 7, _pdf_text(profile.job_title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    meta = (
        f"Location: {profile.location}   |   Reports to: {profile.reports_to}   |   "
        f"Subordinates: {profile.direct_reports}   |   Level: {_grade(profile)}   |   "
        f"Status: {profile.status} (v{profile.version})"
    )
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 5, _pdf_text(meta), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    def section(heading: str, body: str):
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_fill_color(240, 240, 240)
        pdf.set_x(pdf.l_margin)
        pdf.cell(0, 7, _pdf_text(heading), fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 4.6, _pdf_text((body or "—")[:4000]), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.5)

    for heading, attr in TEMPLATE_SECTIONS + REQUIREMENT_SECTIONS:
        section(heading, getattr(profile, attr))

    tech = [pc for pc in profile.profile_competencies if pc.competency and pc.competency.comp_type == "Technical"]
    section("Technical Competencies", "\n".join(f"- {pc.competency.name} ({pc.competency.category}) - {pc.required_level}" for pc in tech))
    behavioral = [pc.competency.name for pc in profile.profile_competencies if pc.competency and pc.competency.comp_type == "Behavioral"]
    section("Behavioral Competencies", "\n".join(f"- {b}" for b in behavioral) or profile.behavioral_text)
    for heading, attr in TAIL_SECTIONS:
        section(heading, getattr(profile, attr))
    section("Signoff", "Line Manager | Functional Manager | OA | HoD | CHRO (if required) | CEO/COO (if required)")

    return bytes(pdf.output())


# ------------------------------------------------------------------ Excel
EXCEL_COLUMNS = [
    ("Job Title", "job_title"),
    ("Job Code", "job_code"),
    ("Division", "division"),
    ("Department", "department"),
    ("Location", "location"),
    ("Reports To", "reports_to"),
    ("Subordinates", "direct_reports"),
    ("Level", "level"),
    ("Role Family", "role_family"),
    ("Job Family", "job_family"),
    ("Status", "status"),
    ("Version", "version"),
    ("Mission", "mission"),
    ("Context", "context"),
    ("Education", "education"),
    ("Experience", "experience"),
    ("Trainings", "trainings"),
    ("Working Conditions", "working_conditions"),
    ("Performance Standards", "performance_standards"),
]


def profiles_to_xlsx(profiles: list[models.JobProfile]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job Profiles"
    ws.append([label for label, _ in EXCEL_COLUMNS])
    for profile in profiles:
        data = profile_to_dict(profile)
        ws.append([str(data.get(key, ""))[:32000] for _, key in EXCEL_COLUMNS])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def import_template_xlsx() -> bytes:
    """Blank bulk-upload template matching the master workbook columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job_Profile_Data"
    columns = [
        "Title of Position", "Location of the Job", "Reports to", "Number of Subordinates",
        "Level", "Mission", "Context", "Roles & Responsibilities", "Collaboration",
        "Authorities", "Education", "Experience", "Technical Competencies",
        "Behavioral Competencies", "Trainings", "General Working Condition",
        "Performance Standards", "Division Code",
    ]
    ws.append(columns)
    ws.append([
        "Example Senior Specialist", "NWG - Transmission Planning", "Example Manager", "0", "3",
        "One-paragraph mission statement...", "Business/environment context...",
        "Statement one. Statement two. Statement three.",
        "Direct Reports: None Matrix Reports to: None Key Customers: ... Key Suppliers: ...",
        "As per Delegation of authority", "BSc Degree in ...", "Minimum of 5 years ...",
        "Network Planning and Design Specialized Knowledge Advanced English Literacy General Business Knowledge Intermediate",
        "Lead with Care\nCollaborate with Agility\nAct with Inclusion\nServe with Respect\nCan-do with Integrity",
        "Relevant technical courses", "General MTNIrancell working conditions",
        "As per performance agreement", "NWG",
    ])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
